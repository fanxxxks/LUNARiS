from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'可重构月球科研站_巨构框架方案_V4.docx'
QA=Path(__file__).resolve().parent
QA.mkdir(exist_ok=True)
FONT='C:/Windows/Fonts/msyh.ttc'

def diagram():
    im=Image.new('RGB',(1600,590),'#f4f7fa'); dr=ImageDraw.Draw(im)
    def ft(s): return ImageFont.truetype(FONT,s)
    levels=[480,387,294,201,108]
    lefts=[205,205,285,285,365]; rights=[1440,1440,1440,1320,1200]
    for i,y in enumerate(levels):
        dr.text((35,y-18),f'L{i+1}',font=ft(25),fill='#243a4c')
        dr.text((91,y-14),f'{8.5+i*27:.1f}m',font=ft(21),fill='#526372')
        dr.line((lefts[i],y,815,y),fill='#637889',width=7)
        dr.line((980,y,rights[i],y),fill='#637889',width=7)
    for x in [390,710,1170]:
        track_top=175 if x==710 else 80
        dr.line((x-10,track_top,x-10,485),fill='#1e748c',width=3)
        dr.line((x+10,track_top,x+10,485),fill='#1e748c',width=3)
    for x,y2 in [(215,480),(300,201),(375,108),(1220,201),(1440,387)]:
        dr.line((x,y2,x,535),fill='#637889',width=8)
    dr.line((185,535,1460,535),fill='#637889',width=5)
    for a,b in [(215,300),(300,375),(1220,1440)]:
        dr.line((a,530,b,480),fill='#9aabb8',width=3)
    positions=[(460,440),(630,440),(1260,440),(260,347),(470,347),(1030,347),(1250,347),(490,254),(650,254),(1280,254),(480,161),(1050,161),(480,68),(990,68)]
    for x,y in positions:
        dr.rounded_rectangle((x,y,x+120,y+35),radius=4,fill='#d5dfe6',outline='#435c6f',width=2)
        dr.rectangle((x+11,y+9,x+73,y+20),fill='#446577')
        dr.line((x+88,y+7,x+109,y+7),fill='#ba804c',width=3)
    dr.rounded_rectangle((834,149,963,469),radius=2,outline='#c0cad2',width=2)
    dr.text((854,242),'贯通',font=ft(25),fill='#697d8e')
    dr.text((854,279),'空腔',font=ft(25),fill='#697d8e')
    for x,name in [(390,'A'),(710,'B'),(1170,'C')]:
        dr.text((x-7,135 if x==710 else 28),name,font=ft(25),fill='#1e748c')
    dr.text((204,551),'框架剖面关系示意  竖向轨道嵌入主框架  非泊位拓扑或施工图',font=ft(21),fill='#5d6e7d')
    p=QA/'frame_section.png'; im.save(p); return p

d=Document(); s=d.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.75);s.bottom_margin=Cm(1.65);s.left_margin=Cm(2);s.right_margin=Cm(2)
s.header_distance=Cm(.65);s.footer_distance=Cm(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=d.styles[name];st.font.name='Microsoft YaHei';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(6)
    for border in st._element.xpath('.//w:pBdr'): border.getparent().remove(border)
d.styles['Normal'].font.size=Pt(10.2);d.styles['Normal'].paragraph_format.line_spacing=1.14
d.styles['Title'].font.size=Pt(25);d.styles['Title'].paragraph_format.space_after=Pt(10)
d.styles['Subtitle'].font.size=Pt(11);d.styles['Subtitle'].font.italic=False
d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 1'].paragraph_format.space_before=Pt(0);d.styles['Heading 1'].paragraph_format.space_after=Pt(12)
d.styles['Heading 2'].font.size=Pt(11.6);d.styles['Heading 2'].paragraph_format.space_before=Pt(10);d.styles['Heading 2'].paragraph_format.space_after=Pt(5)
d.styles['Caption'].font.size=Pt(8.5);d.styles['Caption'].paragraph_format.space_after=Pt(7)
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('可重构月球科研站  V4  ·  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('69737C')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(t,boldlead=None):
    pa=d.add_paragraph()
    if boldlead and t.startswith(boldlead): pa.add_run(boldlead).bold=True;pa.add_run(t[len(boldlead):])
    else:pa.add_run(t)
    return pa
def h(t):d.add_heading(t,2)
def page(t):d.add_page_break();d.add_heading(t,1)
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,x in zip(t.rows[0].cells,headers):c.text=x
    for data in rows:
        for c,x in zip(t.add_row().cells,data):c.text=str(x)
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for ri,row in enumerate(t.rows):
        trp=row._tr.get_or_add_trPr(); ns=OxmlElement('w:cantSplit');trp.append(ns)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            pr.append(b);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'263E50' if ri==0 else ('F2F5F8' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            for pa in c.paragraphs:
                pa.paragraph_format.space_after=Pt(2);pa.paragraph_format.space_before=Pt(2);pa.paragraph_format.line_spacing=1.08
                if ci==0 and widths[ci]<3.5:pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in pa.runs:r.font.size=Pt(9);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '172630')
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

d.add_heading('可重构月球科研站',0)
d.add_paragraph('巨构框架方案与验证路线  V4  2026年9月7日',style='Subtitle')
p('本方案面向课程评审、机械设计讨论与项目路演，提出一座具有错层、退台和局部悬挑的不规则巨型框架。24个功能房间嵌入5层轨道网络，沿框架内部的X、Y、Z方向有序迁移，利用预留泊位完成科研协作、维修隔离和扩建重组。')
p('设计的核心是让框架承担空间组织，让房间承担可替换的任务功能。中央贯通空腔、组团之间的净距与空闲泊位共同保证框架可见、路径可读、调度有余量。建筑的规模通过结构跨度、层高和细部比例呈现。')
h('确定的概念参数')
table(['项目','V4设定','解释'],[
('组织规模','5层  42个有效节点  24个房间','18个节点初始空闲，约42.9%；空闲节点可承担等待与调度，不等同于永久交通面积。'),
('水平间距','X方向31 m  Z方向30 m','最大运输包络为X 23 m × Y 19.5 m × Z 21 m；相邻同向包络间净距约8 m和9 m。'),
('垂直层距','27 m','泊位基准Y为8.5、35.5、62.5、89.5、116.5 m。基准并非室内净高。'),
('竖向迁移','A  B  C三条嵌入式轨道','房间先沿水平轨道到达连通竖向节点，再由托架跨层运输。'),
('调度规则','同一时刻运输一个房间','其他房间保持停靠；空闲泊位用于周转，支撑和封闭通道构成路径约束。')],[3.1,5.25,8.65])
h('本轮形成的判断')
p('巨构框架比紧密排列的独立楼块更符合可重构建筑的表达：骨架长期使用，功能房间按任务迁移。42个节点提供空间余量，但空位比例本身不保证可达性，仍须逐任务检查连通性、占位和运动包络。')
p('各层节点数／初始房间数依次为L1 14／9、L2 11／6、L3 8／5、L4 6／3、L5 3／1，形成逐层收缩的退台轮廓。图形仿真支持几何、占位与时序检查；承载、竖直驱动、锁定制动及生命保障接口需要进一步设计和实物验证。')

page('一 空间与视觉设计')
h('以不规则框架建立整体轮廓')
p('主框架由不同长度的水平跨和高低不一的立柱组织，形成退台、错层及有限悬挑。主要支撑布置在边界和可识别的结构带，中央保留贯穿五层的大空腔。房间分成若干局部组团，组团之间显露轨道、连接节点与维护空间。')
fig=d.add_paragraph();fig.paragraph_format.keep_with_next=True;pic=fig.add_run().add_picture(str(diagram()),width=Cm(17))
pic._inline.docPr.set('descr','五层退台巨构剖面关系示意，中央留空，A和C竖向轨道贯通五层，B至四层，房间沿两侧框架嵌入。仅示意结构关系。')
d.add_paragraph('图1  结构与轨道的剖面关系示意  实际节点和允许路径以仿真拓扑为准',style='Caption')
h('房间是嵌入骨架的功能单元')
p('统一运输底盘、停靠基准和接口区，使不同房间共享搬运机构；外壳尺寸在最大包络内变化。科研舱、样品处理舱、生活舱及设备舱通过窗带比例、舱门、设备面板和局部内装区分，避免只用不同颜色表示用途。固定的连续载人通道如何在迁移前后重接，另列为工程设计问题。')
h('材质与光影服务于结构阅读')
p('框架采用哑光金属与局部磨损边缘，轨道具有较明确的金属反射，舱体外壳表现细微板缝与粗糙度差异，玻璃保留受控反射及部分内部可见性。连接件、锁定机构、导向轮、隔热层和维护走道构成近景细节。材料表现避免通体镜面或过强自发光。')
p('场景以有方向性的主光形成清晰投影，用月面地形起伏和接触阴影落稳建筑；室内暖光与导向冷光控制在局部。日景、低角度照明与夜间检修表达不同观察目的，辉光不能吞没轨道与支撑。月面背景不采用地球式云雾解释真实环境；艺术化补光与曝光属于展示参数。')
h('观察顺序')
p('首先用总览看清退台轮廓和中央空腔，再用分层观察读取可用泊位，随后以房间近景观察运输底盘和锁定接口。机构视图展示承载关系，任务视图保留移动轨迹与终态，减少装饰元素对运动理解的干扰。')

page('二 搬运机构与接口架构')
h('框架内的三维运动')
p('房间的移动遵循已定义的正交轨道图：X和Z方向承担层内平移，Y方向承担跨层升降。A轨道位于列−1、行0，C轨道位于列0、行−1，均贯通5层；B轨道位于列1、行0，贯通L1至L4。各层水平轨道与竖向节点相连，房间经过合法节点转接，不能跨越无连接的空腔或框架。')
p('水平承载、竖向承载与长期停靠需要分别承担功能。概念架构采用运输底盘与可收放的升降托架，房间到站后将重量传递给泊位支承及机械锁定装置；轨道导向约束姿态。驱动形式需通过载荷、维护与冗余取舍确定，暂不把某一种传动方案视为已验证。')
h('建议的动作状态与互锁')
table(['状态','动作与可见表现','进入下一状态的工程条件'],[
('准备','核对目标、路径、占位和接口状态','目标可用，路线未封闭，任务允许迁移。'),
('接管载荷','托架到位，运输底盘受承托','载荷接管确认后才允许解除泊位锁定。'),
('解锁与断接','显示锁扣打开和接口区隔离','接口断接、机械解锁经传感器确认。'),
('导向运输','沿轨道平移或沿Y轴升降','转接点停稳；姿态、速度、制动状态持续受控。'),
('停靠与交接','定位、锁定、托架卸载或收拢','泊位锁定确认后才允许托架离开。'),
('任务结束','显示实际终态与计划差异','位置一致、占位无冲突；服务接口另行验收。')],[2.7,6.4,7.9])
p('上述互锁是工程深化要求。当前仿真可表达动作顺序与几何状态，并不读取真实载荷、制动器或气密接口传感器。串行运输降低同时运动的协调复杂度，多条竖向轨道提供路线选择；独立备份能力仍取决于网络连通和实际机构设计。')
h('标准接口的四个层面')
p('机械接口规定承载点、定位基准、锁定方向和运输包络；电力与数据接口规定隔离、接地、插接顺序和状态确认；流体与热控接口规定关断、泄漏检测和重接检验；载人通行接口规定停靠后连接、疏散与气密验证。统一外形接口只能作为设计起点，不能代替系统兼容性测试。')

page('三 任务价值与仿真指标')
p('每个演示从一个明确任务出发，先给出目标房间和目的泊位，再显示可执行路径、周转过程与结束状态。这样可以讨论迁移带来的收益，也能看见空位、升降行程和串行运输的代价。')
table(['任务','希望产生的变化','演示重点'],[
('科研协作','将R02、R03、R08调入L2，同层协作房间由0个变为3个。','显示跨层路径、参与模块与总距离。同层聚集是协作便利的代理指标，实际效率需任务数据验证。'),
('维修隔离','R24从顶层移至首层东侧检修位，附近两个普通泊位留空作缓冲；R11补入顶层。','显示故障房间移出、邻位缓冲和替代房间停靠。空间移开不等于气密、污染或电气隔离完成。'),
('基地扩建','激活L2两处预留泊位，使该层房间由6个增至8个，全站仍为24个房间。','显示既有房间如何让位和重组。演示表达空间容量的启用，不代表自动制造或安装了新结构和房间。')],[2.6,6.5,7.9])
h('让操作结果可复查')
table(['指标','统计口径','判断用途'],[
('目标完成','逐房间比较实际泊位与目标泊位','验证任务结束状态，不能只以动画结束判定成功。'),
('搬运次数','房间从一个泊位移至另一泊位计一次','显露空位周转引起的额外搬运。'),
('移动距离','累计载房间路径的X Y Z分段长度','比较任务代价；不直接换算成能耗或成本。'),
('升降次数及行程','按跨层段累计次数和绝对Y位移','确认存在真实Y坐标变化，观察竖向运输负担。'),
('占位与包络','检查节点占用及路径扫掠范围','避免穿过静止房间、支撑、边界或封闭通道。'),
('异常与重放','保存初态、目标、路径和事件','复查无路、急停与恢复；性能结论附版本和设备。')],[3.4,6.9,6.7])
h('不把播放时长当作实测性能')
p('仿真时长受展示速度和阶段时序设定影响。路演可用倍率控制节奏，论证效率时须另行测量实际定位、接管载荷、解锁、运动、锁定与接口重接耗时。比较两种布局时应同时展示初态、约束和终态，避免仅选择最短的一次路线。')

page('四 从桌面样机到框架验证')
h('保留原样机作为可测量的第一步')
p('原方案为800 × 600 × 100 mm底座、4 mm面板、八个60 × 40 × 25 mm轻质模型，研究板下XY机构、升降磁头与等待位换位。重量由面板承担，磁耦合用于牵引；磁头靠近或远离面板属于吸合与释放动作。它与本方案的大尺度三维框架具有不同的承载和驱动问题。')
p('原样机继续验证共享驱动、被动模块和等待位周转的基本逻辑。V4增加的Y轴房间运输需要独立的竖向试验台；原XY样机成功不能外推为巨构楼层运输成功。')
table(['阶段','实物或分析工作','进入下一阶段的证据'],[
('1 平面原理','保留八模块与等待位，测板材、质量、速度、脱耦和释放。','完整换位记录、定位误差、失败原因。建议连续20次；原目标19/20及误差≤5 mm均是目标，不是现有结果。'),
('2 单跨单层','用具有明确承载面的底盘替代仅磁牵引展示，加入泊位锁定。','载荷接管与交接可重复；测挠度、定位和锁定可靠性。'),
('3 双层升降','搭建两层一条竖向轨道，先使用非载人配重模型。','证明导向、制动、限位及断电保持；记录跨层全部状态。'),
('4 多层调度','扩至多泊位与两条以上可选择路径，验证通道封闭和恢复。','实物坐标与仿真一致；失败路径被拒绝，事件可重放。'),
('5 接口与环境','完成载荷预算与结构分析，再开展接口及环境专项验证。','形成机械、电气、热控、防尘、热真空和寿命证据；据此决定下一步适用范围。')],[2.7,6.4,7.9])
h('软件验收与画面验收分别记录')
p('软件检查节点邻接、边界、占位、终态及移动全程包络；测试包含三类预设、手动六向调度、目标占用、无可达路径、竖向通道封闭、暂停与急停恢复。随机用例须保存种子、版本、数量与失败记录，不能只报告通过次数。')
p('画面检查总览轮廓、玻璃与金属差异、支撑投影、近景机构、层间可见性和移动时的穿插；交互检查视角切换、房间选择与界面可读性。帧率或加载时间仅在记录的设备与浏览器条件下成立。')
h('资源安排')
p('团队分配机械、电控、软件、报告与路演职责，允许兼任。采购先做材料规格与数量清单，再以实际询价填写预算。优先投入双层升降与停靠交接验证；大尺度巨构的总质量、建造方式与维护成本在概念阶段保持待评估。')

page('五 路演讲述与答辩口径')
h('八分钟演示安排')
table(['时间','画面与操作','讲述重点'],[
('0:00至0:50','总览框架与中央空腔','任务变化会改变房间之间的位置需求。我们研究在共享结构内调整功能房间的位置。'),
('0:50至1:45','切换立面和分层','5层42个节点容纳24个房间，保留18个初始空闲节点，既看清结构也给运输留出周转空间。'),
('1:45至2:45','近看轨道 托架 锁定装置','房间沿水平轨道到达内部升降位置，接管载荷后解锁，再跨层停靠。竖向运动发生在框架内。'),
('2:45至4:00','运行科研协作','展示目标、路径与终态，解释哪几个房间因协作关系调整位置。'),
('4:00至5:00','运行维修隔离或基地扩建','说明任务变化为何需要空位和让位，并显示额外搬运与升降行程。'),
('5:00至6:00','显示指标与测试记录','几何、占位、时序由软件检查；承载、制动和接口需要实物证据。展示实际记录，未测项保持待验证。'),
('6:00至7:10','展示原XY样机和双层试验路线','原样机保留共享驱动与等待位原理，下一步验证房间真正升降及安全交接。'),
('7:10至8:00','回到完成布局','总结任务价值与当前证据，提出机械与环境方向的具体验证需求。')],[2.7,5.6,8.7])
h('三个容易混淆的答辩问题')
p('为什么需要巨型框架：它提供重复使用的承载、导向与停靠基础，使功能房间能够重组；代价是结构质量、建造复杂度和固定轨道范围。是否优于独立移动平台，需比较任务收益与全系统质量、能耗、可靠性和维护。')
p('是否已经实现三维建筑搬运：软件表达房间在内部轨道上的X、Y、Z位移并检查相应约束。现有XY实物验证水平子问题；双层升降、真实承载和锁定互锁尚需新试验。')
p('能否载人移动或长期月面运行：本阶段没有这类实测证据。生命保障、应急疏散、制动冗余、结构与环境适应性均需专项论证，路演不以舱内人物模型作为载人能力证明。')
p('现场保留仿真、实物原理样机和真实样机视频三类材料，并明确标识。无可达路径时展示拒绝原因；设备故障时停止动作，转为静态图解释，不伪装成成功演示。')

page('六 项目依据与方案改进摘要')
h('保留的项目背景')
p('国家航天局2025年4月25日公开材料描述国际月球科研站具有长期自主运行、短期有人参与、可扩展和可维护等特点，并给出当时基本型及拓展型时间安排。[1] 本项目据此讨论可重构设施的研究价值；国家任务背景不构成对本方案的采用、资金或采购承诺。')
p('NASA关于Artemis月面概念设施迁移的研究评估过通过可分离移动平台搬迁中大型设施，并指出系统与架构风险仍需量化。[2] 因此创新论证集中于共享框架、标准泊位、被动功能房间及任务调度的组合与取舍，避免宣称首次提出月面建筑移动。')
p('NASA月尘缓解指南讨论避免暴露、清除和容忍月尘等策略。[3] 主框架、轨道和接口采用防尘设计是必要研究方向；结构遮蔽、封闭面板或装饰性外壳不能直接证明长期免维护。')
h('相较原方案的实质改进')
table(['原有基础','V4改进','保留的边界'],[
('共享板下XY牵引','扩展为框架内水平轨道与3条竖向轨道的概念架构。','平面原理样机继续独立验收，不外推竖向承载能力。'),
('等待位换位','利用42节点中的18个初始空闲节点组织多层周转。','空位多不等于一定有路径，每项任务仍需规划检查。'),
('建筑模型与分层展示','以退台巨构、中央空腔和嵌入式房间统一空间组织。','剖切、分层和爆炸效果是观察工具，不计入真实位移。'),
('基础换位与维护调位','以科研协作、维修隔离和基地扩建解释迁移目的。','位置变化仅是任务收益的代理，实际工作效率另测。'),
('实物与仿真证据分层','增加双层升降、机械交接、接口与环境的分阶段验证。','未经试验的成功率、成本优势和月面适用性不填成结论。')],[3.45,7.1,6.45])
h('参考资料')
for t,u in [
('[1] 国家航天局 送嫦娥架鹊桥 吴伟仁详解国际月球科研站新动态 2025年4月25日','https://www.cnsa.gov.cn/n6758823/n6758844/n10663475/n10663516/c10668780/content.html'),
('[2] NASA Assessing the Relocation of Artemis Foundational Lunar Surface Concepts 2024年IEEE会议相关论文','https://ntrs.nasa.gov/citations/20230013620'),
('[3] NASA Lunar Dust Mitigation A Guide and Reference NASA TP 20220018746','https://ntrs.nasa.gov/api/citations/20220018746/downloads/TP-20220018746.pdf')]:
    pa=p(t);pa.paragraph_format.space_after=Pt(2)
    pa=p(u);pa.paragraph_format.space_after=Pt(7)
    for r in pa.runs:r.font.size=Pt(8)

for para in d.paragraphs:
    if para.style.name in ['Title','Subtitle','Heading 1','Heading 2']:
        for r in para.runs:r.font.color.rgb=RGBColor(0,0,0)
        for b in para._p.xpath('./w:pPr/w:pBdr'):b.getparent().remove(b)
d.core_properties.title='可重构月球科研站 巨构框架方案与验证路线 V4'
d.core_properties.subject='巨构空间设计 轨道搬运 任务价值 分阶段验证'
d.core_properties.author='';d.core_properties.last_modified_by=''
d.save(OUT)
print(str(OUT))
