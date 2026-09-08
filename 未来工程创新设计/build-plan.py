from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

d=Document()
s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Microsoft YaHei';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.12
d.styles['Title'].font.size=Pt(24);d.styles['Heading 1'].font.size=Pt(18);d.styles['Heading 2'].font.size=Pt(12)
d.styles['Heading 2'].paragraph_format.space_before=Pt(8)
d.styles['Heading 2'].paragraph_format.space_after=Pt(4)
for st in d.styles:
 for border in st._element.xpath('.//w:pBdr'):
  border.getparent().remove(border)
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,2)
def page(t):d.add_page_break();d.add_heading(t,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,x in zip(t.rows[0].cells,headers):c.text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=x
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for ri,row in enumerate(t.rows):
  for ci,c in enumerate(row.cells):
   c.width=Cm(widths[ci]);c.vertical_alignment=1;pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEE9' if ri==0 else 'FFFFFF');pr.append(shade)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(3)
    for r in para.runs:r.font.size=Pt(9);r.bold=ri==0

d.add_heading('可重构月球科研站',0)
p('项目策划与路演方案  修订版  2026年9月7日')
h('一 项目定位')
p('我们研究月面科研设施随任务变化进行模块换位的可能性，提出共享式板下 XY 搬运机构，使多个被动模块通过一个等待位完成平面重排。本轮交付包括八舱段原理样机、独立运行的三维仿真软件和可复查的测试记录。')
p('核心主张是把基地布局从设计时变量变成运行时变量。当前样机研究的范围为预设平台上的水平换位；垂直堆叠、载人移动和真实月面部署列入未来研究，不作为本轮已实现能力。')
h('二 应用背景与创新边界')
p('国家航天局2025年公开材料将国际月球科研站描述为长期自主运行、短期有人参与、可扩展、可维护的综合科研设施；当时规划为2035年前基本型、2045年前拓展型。[1] 这提供了研究背景，并不代表本项目已获得任务方需求确认、投资或采购承诺。')
p('既有研究已讨论移动栖居舱和可分离搬运平台。[2] 因此本项目不主张首次提出月面设施移动，而比较共享机构、被动模块和平台基础设施之间的成本与可靠性取舍。')
table(['设计点','潜在收益','对应代价或未决问题'],[
('共享 XY 搬运机构','减少每个舱段自带驱动的需求','串行吞吐与单点失效；须计入平台质量'),
('驱动位于面板下','降低运动部件直接暴露风险','板缝、维护口与密封件仍需防尘设计'),
('等待位换位','使被占用目标之间的换位可执行','增加搬运次数与预留占地'),
('模块位置可配置','适应科研协作或维护任务','接口、管线、遮挡与承载尚未验证')],[3.5,5.4,8.1])
h('三 课程约束与交付口径')
p('若课程要求建筑真正进行三维搬运或垂直堆叠，现有 XY 样机只能覆盖平面子问题；应在验收前明确这一差距，不能用三维画面或爆炸视图代替 Z 轴搬运。真实买方要求则需补充具体机构访谈、公开征集或项目需求证据。')

page('四 样机与仿真规格')
p('实物外形以原三维 HTML 为参照：800×600×100 mm 底座、4 mm 面板、八个60×40×25 mm轻质舱段。材料和厚度最终以牵引试验及加工结果确定。')
table(['部件','本轮功能','验收证据'],[
('XY 机构与升降磁头','定位、靠近、拖动、远离','实机连续动作录像与位置测量'),
('被动舱段及底部钢片','承受隔板牵引并保持姿态','不同质量、速度下的脱耦记录'),
('面板与等待区','支持舱段并提供换位空间','面板变形、摩擦及释放稳定性'),
('控制与供电','归零、运动、停机和恢复','限位、断电、急停测试记录')],[4,6,7])
h('软件功能')
p('软件提供基础换位、科研协作、维护调位和自定义目标布局。规划器用等待区分解置换循环；每次只移动一个舱段。路径在正交节点图上搜索，考虑舱段矩形外形和安全间距，并可绕开用户启用的中央封闭区。')
p('软件提供暂停、逐阶段播放、时间轴、急停、速度设置、透视、分层、俯视和路演模式，并导出 JSON 计划与运行事件。目标重复或无路径时显示失败原因。每次重新规划均从统一初始布局开始；不支持从搬运中途重规划。')
p('拖动段采用端点速度为零的平滑插值，转角停稳。40、80、120 mm/s为假设牵引峰值；播放倍率仅改变观看速度。预测时长含定位、吸合、分段拖动、释放和空返，不能替代实机测时。')
h('验证边界')
table(['层级','本轮能给出的证据','不可据此推出'],[
('几何与时序仿真','路径、终态、等待区释放、动作顺序','磁力、实际摩擦、热真空与寿命'),
('桌面样机试验','在所测板材、质量与速度下运行','真实尺度载荷或月面环境适用性'),
('应用概念','功能舱用途和目标布局设定','载人安全、气密通道和自动接驳已实现')],[3.4,6.8,6.8])

page('五 执行安排与验收')
table(['阶段','工作与交付','通过条件'],[
('第1周','先完成单舱段磁吸牵引；记录板材、厚度、质量和电源','完成吸合、牵引、释放；遇失败可定位原因'),
('第2周','完成八舱段与等待区；核对仿真和实物坐标','连续完整执行基础换位；无碰撞和手动补位'),
('第3周','重复测试、故障演练、路演和彩排视频','提交全部成功与失败记录；备份素材能独立播放')],[2.4,8.8,5.8])
p('负责人由团队分配为机械、电控、软件、报告与路演四类角色，可兼任。预算按材料规格、数量、单价和供应来源逐项登记；不填未经询价的总额。')
h('实物测试表')
table(['指标','测量方法','本轮建议目标'],[
('完整换位成功率','连续20次，统计无需人工干预的完成次数','目标≥19/20；如未达到，报告真实结果'),
('最终定位误差','每次测中心相对目标中心的最大偏差','目标≤5 mm，结合制作精度调整'),
('牵引质量与速度','逐级增加质量和速度，记录脱耦、偏转','报告可重复运行区间，禁止只报最好一次'),
('释放与空返','观察空返时已放置舱段是否再移动','20次均无明显二次牵引'),
('异常恢复','急停、卡滞和断电后人工确认再恢复','无自行重启，明确归零与重放流程')],[3.8,7.3,5.9])
h('现场执行')
p('机械操作者负责样机；讲解者负责叙事；软件操作者负责预设和画面。开场前确认初始位置P1—P8与R1—R8对应，等待位W为空。先播放基础换位，再展示中央通道封闭后的规划变化，最后回到完整目标布局。')
p('现场保留三个层次：实物原理样机、三维仿真、事先录制的视频与三张静态步骤图。视频必须在真实样机制作后录制；软件画面只能标为仿真，不能冒充实机。故障时暂停设备并切换说明材料。')
h('软件验收')
p('检查基础换位需三次搬运，结束后W为空；检查多个循环与自定义排列；检查重复编号、封闭起点与无路径的拒绝；独立采样检查路径扫掠范围与静止舱段不重叠。浏览器视觉与交互验收在路演设备上另行执行。')

page('六 八分钟路演脚本')
p('以下为讲解词与操作节拍；演示、停顿和切换计入八分钟。实际语速与实机时间以彩排为准。')
h('00:00至00:50 问题')
p('“科研站的任务会变化。新增实验、设备维护和舱段替换，都可能使原来的布局不再合适。我们关注的是：设施建成以后，能否利用同一套搬运设备，让已有模块调整位置？”')
h('00:50至01:40 背景与差异')
p('“已有研究讨论过移动栖居舱和搬运平台。我们的切入点是一个预设的平台：把主要驱动机构集中到面板下方，让上面的模块保持被动。这样可能减少模块自带的运动部件，但也增加了平台建设、串行搬运和维修设计的要求。”')
h('01:40至02:40 拆解原理')
p('操作：开启软件分层，展示XY横梁、滑块、升降磁头和面板。')
p('“下面是一套XY机构，上面是八个轻质模型。磁头移动到目标下面，靠近面板后拖动模块，到位后远离面板。这个升降只负责吸合和释放。画面中的分层用于看清结构，我们本轮没有做舱段垂直堆叠。”')
h('02:40至04:00 基础换位')
p('操作：关闭分层，选择基础任务，规划并播放；样机按实测节奏同步或独立演示。')
p('“R1和R2要交换，但两个目标位都已被占用。我们先将R1放入等待区，再把R2搬到P1，最后将R1搬到P2。一次只搬一个模块，其他模块保持静止。完成后等待区重新空闲，可用于下一次任务。”')
h('04:00至05:10 改变任务')
p('操作：选择科研协作或维护调位，展示目标选择和计划列表；启用中央封闭区后重新规划。')
p('“这里可以改变目标，而不仅是重复一段录像。软件会重新计算搬运顺序和路径。红色区域代表不可通过的区域。如果找不到满足尺寸约束的路径，系统应明确拒绝，而不是让模型穿过障碍。”')
h('05:10至06:20 证据')
p('“我们把证据分成两个层次。软件检查几何路径、顺序和目标状态；样机测量定位误差、成功率、牵引质量和释放效果。”展示本团队实测表；未测项目明确说“尚未测量”，不得读出建议目标作为结果。')
h('06:20至07:15 局限与后续')
p('“面板下方的布置可能减少直接暴露，但板缝和维护口仍然需要防尘。月面应用还需要解决承载、接口、温度变化和维修。我们会先验证桌面机制，再确定是否值得推进环境试验。载人移动与垂直搬运不在本轮成果范围内。”')
h('07:15至08:00 价值与请求')
p('“本项目把布局重构变成一个可以操作和测量的问题。下一步希望获得机械结构与空间环境方向的指导，评估共享平台在质量、可靠性和任务收益上是否有优势。本轮提交的是原理样机、仿真软件和测试记录。”')

page('七 答辩问题与风险')
table(['可能的问题','建议回答'],[
('创新在哪里','既有研究已涉及舱段移动。我们研究共享板下驱动与被动模块的组合及其取舍，尚不声称首创或优于全部方案。'),
('为何称华容道','借用等待位周转的直观类比。单机构并不在数学上必然等同于传统华容道；我们的通道图和运动规则另行定义。'),
('十年不进尘如何保证','当前没有这项证明。板下布置是降低暴露的设计方向；接缝、维护口和密封寿命均需测试。'),
('磁吸是否承重','样机重量由面板支撑，磁耦合提供牵引。真实舱段的水平力传递、锁定和承载机构尚未设计验证。'),
('有人时能否移动','本轮不验证载人移动，也不据活动分级宣称安全。接口断接、应急疏散和动态载荷需要独立研究。'),
('单机构坏了怎么办','它会阻塞后续搬运，是明确的单点失效。实际系统需可维护设计、备件或冗余，并计入质量预算。'),
('谁会买','潜在研究需求方包括相关科研团队与任务论证单位；目前未获得真实采购或投资承诺，下一步通过访谈与公开需求核实。'),
('为什么更省钱','目前没有成本结论。应同时比较驱动数量、整个平台、发射质量、安装与维护开销。'),
('如何实现防尘和供能','本轮未验证。正压吹气涉及气体消耗，太阳能铺设与搬运区可能冲突，均不能作为现成答案。')],[4.1,12.9])
h('风险处理原则')
p('实物牵引不可靠时先降低质量或速度并重新测量；无法通过就如实报告限制。软件无路径时保留失败状态并解释原因。路演不以“对手都不可行”代替数据，不以国家任务存在代替客户证据。')

page('八 来源与交付清单')
p('以下为修订时核对的公开来源。任务安排和合作数量按各来源的发布日期引用，提交前检查有无更新。')
sources=[
('[1] 国家航天局 2025年 国际月球科研站新动态','https://www.cnsa.gov.cn/n6758823/n6758844/n10663475/n10663516/c10668780/content.html','支持2035年前基本型、2045年前拓展型及当时合作数量；不证明本项目有买方。'),
('[2] NASA 2024 Assessing the Relocation of Artemis Foundational Lunar Surface Concepts','https://ntrs.nasa.gov/citations/20240002265','支持已有月面栖居设施搬迁与可分离移动平台研究，不能宣称所有现有方案布局固定。'),
('[3] NASA Lunar Dust Mitigation A Guide and Reference','https://ntrs.nasa.gov/citations/20220018746','支持避免、清除、容忍的综合取舍；不支持单靠板下布置保证十年零尘。'),
('[4] NASA 2020 Game Changing Development Program 访谈','https://www.nasa.gov/podcasts/small-steps-giant-leaps/small-steps-giant-leaps-episode-42-game-changing-development-program/','访谈提到Schmitt靴子类似凯夫拉材料的磨损。本路演删除了容易误引的精确材料与天数叙述。')]
for title,url,meaning in sources:h(title);p(url);p(meaning)
h('本轮文件')
p('可重构月球科研站项目策划与路演方案_修订版.docx：策划、边界、验收、路演与答辩。')
p('月宫华容_三维仿真软件.html：内嵌三维库与规划器，双击本地打开，不依赖CDN。导出的JSON为软件运行记录。')
p('仿真软件使用说明.md：使用步骤、限制与本机验收方法。源代码及自动检查脚本随工作目录保留。')
h('实物完成后补交')
p('材料与成本明细、连续测试原始记录、实际性能表、实机演示视频与静态步骤图。未经测试的项保持待验证。')
in_script=False
for para in d.paragraphs:
 if para.text=='六 八分钟路演脚本':in_script=True
 if para.text=='七 答辩问题与风险':in_script=False
 if in_script:
  para.paragraph_format.space_before=Pt(5 if para.style.name.startswith('Heading') else 0)
  para.paragraph_format.space_after=Pt(3)
  para.paragraph_format.line_spacing=1.0
 if para.style.name=='Title':
  borders=OxmlElement('w:pBdr')
  for side in ['top','bottom','left','right','between']:
   e=OxmlElement('w:'+side);e.set(qn('w:val'),'nil');borders.append(e)
  para._p.get_or_add_pPr().append(borders)
d.save('可重构月球科研站项目策划与路演方案_修订版.docx')
print('Saved revised DOCX')
